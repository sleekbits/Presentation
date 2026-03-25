from pathlib import Path

from pptx import Presentation

from backend.services.pptx_parser import extract_placeholders


def build_template(path: Path) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    textbox = slide.shapes.add_textbox(0, 0, 6000000, 1000000)
    tf = textbox.text_frame
    p = tf.paragraphs[0]
    p.add_run().text = "Award {{pro"
    p.add_run().text = "jectTitle}} for <<clientName>>"
    table = slide.shapes.add_table(1, 1, 0, 1500000, 6000000, 1000000).table
    table.cell(0, 0).text = "Budget [[budget]]"
    prs.save(str(path))


def test_extract_placeholders(tmp_path: Path) -> None:
    template = tmp_path / "template.pptx"
    build_template(template)

    result = extract_placeholders(str(template))
    keys = {p["key"] for p in result["placeholders"]}
    assert "projectTitle" in keys
    assert "clientName" in keys
    assert "budget" in keys
