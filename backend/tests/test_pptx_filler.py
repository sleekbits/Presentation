from pathlib import Path

from pptx import Presentation

from backend.services.pptx_filler import fill_pptx


def build_template(path: Path) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    textbox = slide.shapes.add_textbox(0, 0, 6000000, 1000000)
    textbox.text_frame.paragraphs[0].text = "Hello {{projectTitle}}"
    table = slide.shapes.add_table(1, 1, 0, 1500000, 6000000, 1000000).table
    table.cell(0, 0).text = "Budget: [[budget]]"
    prs.save(str(path))


def test_fill_replaces_text_and_table(tmp_path: Path) -> None:
    template = tmp_path / "template.pptx"
    output = tmp_path / "out.pptx"
    build_template(template)

    fill_pptx(template, output, {"projectTitle": "Bridge Upgrade", "budget": "AED 10M"})

    prs = Presentation(str(output))
    slide = prs.slides[0]
    textbox_text = slide.shapes[0].text_frame.text
    table_text = slide.shapes[1].table.cell(0, 0).text
    assert "Bridge Upgrade" in textbox_text
    assert "AED 10M" in table_text
