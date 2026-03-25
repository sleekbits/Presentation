from pathlib import Path

from fastapi import APIRouter, Depends, HTTPException
from pptx import Presentation

from ..schemas.api import UploadResponse
from ..services.job_store import job_store
from ..utils.auth import require_token
from ..utils.config import settings

router = APIRouter(prefix="/api", tags=["sample"])


SAMPLE_PATH = Path(settings.temp_dir) / "sample_template.pptx"


def _ensure_sample() -> None:
    if SAMPLE_PATH.exists():
        return
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Award Recommendation {{projectTitle}}"
    slide.placeholders[1].text = "Client: {{clientName}} | Budget: {{budget}}"
    slide2 = prs.slides.add_slide(prs.slide_layouts[5])
    table = slide2.shapes.add_table(2, 2, 0, 0, 6000000, 1500000).table
    table.cell(0, 0).text = "Duration"
    table.cell(0, 1).text = "{{duration}}"
    table.cell(1, 0).text = "Start Date"
    table.cell(1, 1).text = "[[startDate]]"
    prs.save(str(SAMPLE_PATH))


@router.post("/sample", response_model=UploadResponse, dependencies=[Depends(require_token)])
def use_sample_template() -> UploadResponse:
    _ensure_sample()
    if not SAMPLE_PATH.exists():
        raise HTTPException(status_code=500, detail="Failed to build sample")
    destination = Path(settings.temp_dir) / "incoming" / f"sample-{SAMPLE_PATH.name}"
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_bytes(SAMPLE_PATH.read_bytes())
    job = job_store.create_job(destination)
    job.add_audit("upload", filename="sample_template.pptx", size=destination.stat().st_size)
    return UploadResponse(jobId=job.id, state=job.state.value)
