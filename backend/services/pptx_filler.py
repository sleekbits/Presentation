from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from pptx import Presentation

TOKEN_REGEX = re.compile(r"(\{\{\s*([a-zA-Z0-9_.-]+)\s*\}\}|\[\[\s*([a-zA-Z0-9_.-]+)\s*\]\]|<<\s*([a-zA-Z0-9_.-]+)\s*>>)" )


def _resolve_key(match: re.Match[str]) -> str:
    return next(group for group in match.groups()[1:] if group)


def _replace_tokens(text: str, answers: dict[str, Any], keep_unfilled: bool) -> str:
    def repl(match: re.Match[str]) -> str:
        key = _resolve_key(match)
        value = answers.get(key)
        if value is None or value == "":
            return match.group(0) if keep_unfilled else ""
        return str(value)

    return TOKEN_REGEX.sub(repl, text)


def _replace_paragraph(paragraph, answers: dict[str, Any], keep_unfilled: bool) -> None:
    full = "".join(run.text for run in paragraph.runs)
    if not full:
        return
    replaced = _replace_tokens(full, answers, keep_unfilled)
    if replaced == full:
        return
    if paragraph.runs:
        paragraph.runs[0].text = replaced
        for run in paragraph.runs[1:]:
            run.text = ""


def fill_pptx(template_path: Path, output_path: Path, answers: dict[str, Any], keep_unfilled: bool = False, include_notes: bool = False) -> None:
    prs = Presentation(str(template_path))

    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text_frame") and shape.text_frame is not None:
                for paragraph in shape.text_frame.paragraphs:
                    _replace_paragraph(paragraph, answers, keep_unfilled)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    for cell in row.cells:
                        if cell.text_frame is None:
                            continue
                        for paragraph in cell.text_frame.paragraphs:
                            _replace_paragraph(paragraph, answers, keep_unfilled)
        if include_notes and slide.has_notes_slide:
            for shape in slide.notes_slide.shapes:
                if hasattr(shape, "text_frame") and shape.text_frame is not None:
                    for paragraph in shape.text_frame.paragraphs:
                        _replace_paragraph(paragraph, answers, keep_unfilled)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))
