from __future__ import annotations

import re
from collections import OrderedDict

from pptx import Presentation

PLACEHOLDER_PATTERNS = [
    ("{{}}", re.compile(r"\{\{\s*([a-zA-Z0-9_.-]+)\s*\}\}")),
    ("[[]]", re.compile(r"\[\[\s*([a-zA-Z0-9_.-]+)\s*\]\]")),
    ("<<>>", re.compile(r"<<\s*([a-zA-Z0-9_.-]+)\s*>>")),
]


def infer_type(key: str, context: str) -> str:
    k = key.lower()
    c = context.lower()
    if any(v in k for v in ["date", "deadline", "due"]):
        return "date"
    if any(v in k for v in ["budget", "price", "cost", "amount", "aed", "usd"]) or "aed" in c:
        return "currency"
    if any(v in k for v in ["count", "qty", "number", "score"]):
        return "number"
    if any(v in k for v in ["description", "summary", "notes", "scope", "background"]):
        return "longtext"
    return "string"


def _iter_text_blocks(slide, include_notes: bool):
    for shape in slide.shapes:
        shape_name = getattr(shape, "name", None)
        if hasattr(shape, "text_frame") and shape.text_frame is not None:
            for p_idx, paragraph in enumerate(shape.text_frame.paragraphs):
                yield "paragraph", shape_name, p_idx, "".join(run.text for run in paragraph.runs)
        if getattr(shape, "has_table", False):
            for r_idx, row in enumerate(shape.table.rows):
                for c_idx, cell in enumerate(row.cells):
                    yield "table", shape_name, f"{r_idx}:{c_idx}", cell.text

    if include_notes and slide.has_notes_slide:
        for shape in slide.notes_slide.shapes:
            shape_name = getattr(shape, "name", None)
            if hasattr(shape, "text_frame") and shape.text_frame is not None:
                for p_idx, paragraph in enumerate(shape.text_frame.paragraphs):
                    yield "notes", shape_name, p_idx, "".join(run.text for run in paragraph.runs)


def extract_placeholders(path: str, include_notes: bool = False) -> dict:
    prs = Presentation(path)
    found: OrderedDict[tuple[str, str], dict] = OrderedDict()

    for slide_index, slide in enumerate(prs.slides, start=1):
        for _kind, shape_name, _loc, text in _iter_text_blocks(slide, include_notes):
            if not text:
                continue
            for fmt, pattern in PLACEHOLDER_PATTERNS:
                for match in pattern.finditer(text):
                    key = match.group(1)
                    if not re.fullmatch(r"[a-zA-Z0-9_.-]+", key):
                        continue
                    bucket_key = (key, fmt)
                    if bucket_key not in found:
                        found[bucket_key] = {
                            "key": key,
                            "format": fmt,
                            "slideIndex": slide_index,
                            "shapeName": shape_name,
                            "context": text[:180],
                            "suggestedType": infer_type(key, text),
                        }

    return {"placeholders": list(found.values())}
